import csv
import io
import json
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from freedfe.documentos.leitor import LeitorDocumentoFiscal
from freedfe.excecoes import ValidacaoErro
from freedfe.exportacao.base import COLUNAS
from freedfe.exportacao.fabrica import FabricaExportador
from freedfe.exportacao.formato import FormatoSaida

from fabrica_xml import CNPJ_FORNECEDOR, gerar_chave, nfe_proc


@pytest.fixture
def resumos():
    leitor = LeitorDocumentoFiscal()
    return [leitor.ler(nfe_proc(gerar_chave("55", CNPJ_FORNECEDOR, n))) for n in (1, 2)]


def test_csv_padrao_pt_br(resumos):
    conteudo = FabricaExportador().obter("csv").exportar(resumos)
    linhas = list(csv.DictReader(io.StringIO(conteudo), delimiter=";"))
    assert list(linhas[0]) == list(COLUNAS)
    assert linhas[0]["valor_total"] == "1500,50" and linhas[0]["xml_completo"] == "S"


def test_json(resumos):
    dados = json.loads(FabricaExportador().obter(FormatoSaida.JSON).exportar(resumos))
    assert len(dados) == 2 and dados[0]["valor_total"] == 1500.5
    assert dados[0]["data_emissao"] == "2026-07-15T10:30:00-03:00"


def test_xml(resumos):
    raiz = ET.fromstring(FabricaExportador().obter("XML").exportar(resumos))
    assert raiz.get("quantidade") == "2"
    assert raiz.find("documento/valor_total").text == "1500.50"


def test_xlsx(resumos, tmp_path):
    destino = FabricaExportador().obter("xlsx").gravar(resumos, tmp_path / "saida.xlsx")
    aba = openpyxl.load_workbook(destino).active
    assert [c.value for c in aba[1]] == list(COLUNAS)
    assert aba.max_row == 3
    assert aba.cell(2, COLUNAS.index("valor_total") + 1).value == 1500.5


def test_csv_gravado_com_bom(resumos, tmp_path):
    destino = FabricaExportador().obter("csv").gravar(resumos, tmp_path / "saida.csv")
    assert destino.read_bytes().startswith(b"\xef\xbb\xbf")


def test_formato_invalido():
    with pytest.raises(ValidacaoErro):
        FormatoSaida.normalizar("pdf")
