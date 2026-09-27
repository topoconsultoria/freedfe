"""Exportação em XLSX (requer o extra ``xlsx``: ``pip install freedfe[xlsx]``)."""

from __future__ import annotations

import io
from collections.abc import Sequence

from freedfe.documentos.resumo import ResumoDocumento
from freedfe.excecoes import DependenciaAusenteErro
from freedfe.exportacao.base import COLUNAS, Exportador, linha

FORMATO_VALOR = "#,##0.00"
FORMATO_DATA = "dd/mm/yyyy hh:mm:ss"


class ExportadorXlsx(Exportador):
    """Planilha com cabeçalho congelado, filtro automático e tipos nativos (data e número)."""

    extensao = "xlsx"

    def __init__(self, nome_aba: str = "DFe"):
        self._nome_aba = nome_aba

    def exportar(self, resumos: Sequence[ResumoDocumento]) -> bytes:
        openpyxl = _importar_openpyxl()
        pasta = openpyxl.Workbook()
        aba = pasta.active
        aba.title = self._nome_aba
        aba.append(list(COLUNAS))
        for celula in aba[1]:
            celula.font = openpyxl.styles.Font(bold=True)
        for resumo in resumos:
            aba.append([self._valor(coluna, valor) for coluna, valor in linha(resumo).items()])
        self._formatar_colunas(aba)
        saida = io.BytesIO()
        pasta.save(saida)
        return saida.getvalue()

    @staticmethod
    def _valor(coluna: str, valor):
        if coluna == "data_emissao" and valor is not None:
            return valor.replace(tzinfo=None)  # Excel não armazena fuso; mantém hora local do XML
        if coluna == "xml_completo":
            return "S" if valor else "N"
        return valor

    @staticmethod
    def _formatar_colunas(aba) -> None:
        aba.freeze_panes = "A2"
        aba.auto_filter.ref = aba.dimensions
        indice_valor = COLUNAS.index("valor_total") + 1
        indice_data = COLUNAS.index("data_emissao") + 1
        for (celula_data, celula_valor) in zip(
                aba.iter_rows(min_row=2, min_col=indice_data, max_col=indice_data),
                aba.iter_rows(min_row=2, min_col=indice_valor, max_col=indice_valor)):
            celula_data[0].number_format = FORMATO_DATA
            celula_valor[0].number_format = FORMATO_VALOR
        aba.column_dimensions["A"].width = 48


def _importar_openpyxl():
    try:
        import openpyxl
        import openpyxl.styles  # noqa: F401
    except ImportError as erro:
        raise DependenciaAusenteErro(
            "Exportação XLSX requer openpyxl: pip install 'freedfe[xlsx]'.") from erro
    return openpyxl
